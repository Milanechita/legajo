"""Entrada y salida en los formatos que el area ya usa.

El equipo de cumplimiento trabaja en Excel. Una herramienta que lo obligue a
abandonarlo es la solucion teoricamente correcta e inutil en la practica.
Entra CSV o XLSX, sale XLSX. La automatizacion pasa por debajo; el analista
no cambia de entorno.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from .alertas import Alerta
    from .beneficiario import Resolucion
    from .riesgo import Evaluacion

from .modelo import Caso, Cliente, Documento
from .operaciones import Operacion, Perfil
from .paises import iso
from .pep import PEP, RegistroPEP
from .screening import Coincidencia, ResultadoScreening
from .societaria import Administracion, Control, Estructura, Nodo, Participacion

FUENTE = "Arial"

_ENCABEZADO = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
_RELLENO_ENCABEZADO = PatternFill("solid", start_color="1F3864")
_CUERPO = Font(name=FUENTE, size=10)
_PROBABLE = PatternFill("solid", start_color="F8CBAD")
_REVISION = PatternFill("solid", start_color="FFF2CC")

COLUMNAS_PADRON = [
    "cliente_id", "nombre", "tipo", "documento_tipo", "documento_numero",
    "fecha_nacimiento", "nacionalidad", "pais_residencia", "actividad",
    "oferta_publica",
]


def _booleano(valor: str) -> bool:
    return (valor or "").strip().lower() in {"si", "sí", "s", "true", "1", "x", "yes"}


def _fila_a_cliente(fila: dict[str, str]) -> Cliente:
    documentos = []
    tipo_doc = (fila.get("documento_tipo") or "").strip()
    num_doc = (fila.get("documento_numero") or "").strip()
    if tipo_doc and num_doc:
        documentos.append(Documento(tipo=tipo_doc, numero=num_doc))

    def opcional(clave: str) -> str | None:
        valor = (fila.get(clave) or "").strip()
        return valor or None

    return Cliente(
        cliente_id=(fila.get("cliente_id") or "").strip(),
        nombre=(fila.get("nombre") or "").strip(),
        tipo=(fila.get("tipo") or "PERSONA").strip().upper() or "PERSONA",
        documentos=documentos,
        fecha_nacimiento=opcional("fecha_nacimiento"),
        nacionalidad=opcional("nacionalidad"),
        pais_residencia=opcional("pais_residencia"),
        actividad=opcional("actividad"),
        oferta_publica=_booleano(fila.get("oferta_publica", "")),
    )


def _leer_tabla(ruta: Path) -> list[dict[str, str]]:
    """Lee CSV o XLSX a una lista de diccionarios. Un solo camino para ambos."""
    if ruta.suffix.lower() in {".xlsx", ".xlsm"}:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro.active
        filas = hoja.iter_rows(values_only=True)
        encabezados = [str(c or "").strip() for c in next(filas)]
        registros = [
            dict(zip(encabezados, [("" if v is None else str(v)) for v in fila]))
            for fila in filas
        ]
        libro.close()
        return registros

    with ruta.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def leer_padron(ruta: str | Path) -> list[Cliente]:
    """Lee el padron de clientes desde CSV o XLSX."""
    clientes = [_fila_a_cliente(r) for r in _leer_tabla(Path(ruta))]
    return [c for c in clientes if c.cliente_id and c.nombre]


COLUMNAS_ESTRUCTURA = [
    "relacion", "origen_id", "origen_nombre", "origen_tipo", "origen_jurisdiccion",
    "origen_oferta_publica", "destino_id", "capital", "voto", "detalle",
]

COLUMNAS_PEP = ["cliente_id", "tipo", "cargo", "fecha_cese", "por_parentesco"]


def leer_peps(ruta: str | Path) -> RegistroPEP:
    """Lee las declaraciones juradas de condicion PEP.

    Una condicion cesada hace mas de dos anios ya no es PEP, pero el archivo
    la conserva igual: el dato sirve para el expediente, y borrarlo obligaria
    a volver a pedirle la declaracion al cliente.
    """
    from datetime import datetime

    peps: list[PEP] = []
    for fila in _leer_tabla(Path(ruta)):
        cliente_id = (fila.get("cliente_id") or "").strip()
        if not cliente_id:
            continue

        cese = None
        crudo = (fila.get("fecha_cese") or "").strip()
        if crudo:
            try:
                cese = datetime.fromisoformat(crudo.split("T")[0]).date()
            except ValueError:
                cese = None

        peps.append(PEP(
            cliente_id=cliente_id,
            tipo=(fila.get("tipo") or "NACIONAL").strip().upper() or "NACIONAL",
            cargo=(fila.get("cargo") or "").strip(),
            fecha_cese=cese,
            por_parentesco=_booleano(fila.get("por_parentesco", "")),
        ))

    return RegistroPEP(peps)


def _fraccion(valor: str) -> float:
    """Acepta 0.60, 60 y 60% como la misma cosa."""
    texto = (valor or "").strip().replace("%", "").replace(",", ".")
    if not texto:
        return 0.0
    numero = float(texto)
    return numero / 100.0 if numero > 1.0 else numero


def leer_estructura(ruta: str | Path, clientes: list[Cliente] | None = None) -> Estructura:
    """Lee el grafo societario.

    Una sola tabla con una columna `relacion` que despacha el tipo. Agregar un
    tipo de vinculo nuevo es agregar una rama en el despacho, no un archivo
    de entrada nuevo ni un formato paralelo.
    """
    estructura = Estructura()

    for cliente in clientes or []:
        estructura.agregar_nodo(Nodo(
            id=cliente.cliente_id,
            nombre=cliente.nombre,
            tipo=cliente.tipo,
            jurisdiccion=(cliente.pais_residencia or "AR").strip().upper() or "AR",
            oferta_publica=cliente.oferta_publica,
        ))

    for fila in _leer_tabla(Path(ruta)):
        relacion = (fila.get("relacion") or "").strip().upper()
        origen = (fila.get("origen_id") or "").strip()
        destino = (fila.get("destino_id") or "").strip()
        if not relacion or not origen or not destino:
            continue

        estructura.agregar_nodo(Nodo(
            id=origen,
            nombre=(fila.get("origen_nombre") or origen).strip(),
            tipo=(fila.get("origen_tipo") or "PERSONA").strip().upper() or "PERSONA",
            jurisdiccion=(fila.get("origen_jurisdiccion") or "AR").strip().upper() or "AR",
            oferta_publica=_booleano(fila.get("origen_oferta_publica", "")),
        ))

        detalle = (fila.get("detalle") or "").strip()

        if relacion == "PARTICIPACION":
            capital = _fraccion(fila.get("capital", ""))
            voto_crudo = (fila.get("voto") or "").strip()
            estructura.agregar_participacion(Participacion(
                propietario=origen,
                participada=destino,
                capital=capital,
                voto=_fraccion(voto_crudo) if voto_crudo else None,
            ))
        elif relacion == "CONTROL":
            estructura.agregar_control(Control(
                persona=origen, entidad=destino,
                motivo=detalle or "control final por otros medios",
            ))
        elif relacion == "ADMINISTRACION":
            estructura.agregar_administracion(Administracion(
                persona=origen, entidad=destino,
                cargo=detalle or "administrador",
            ))

    return estructura


def _ajustar_anchos(hoja, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def _escribir_encabezado(hoja, columnas: list[str]) -> None:
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.freeze_panes = "A2"


def exportar(
    resultado: ResultadoScreening,
    casos: list[Caso],
    ruta: str | Path,
    umbral_probable: float,
) -> Path:
    """Escribe el informe de screening en XLSX.

    Tres hojas: resumen, coincidencias, y expediente completo. La tercera es
    la que importa en una inspeccion.
    """
    ruta = Path(ruta)
    libro = Workbook()

    # --- Hoja 1: resumen ---
    resumen = libro.active
    resumen.title = "Resumen"
    _escribir_encabezado(resumen, ["Concepto", "Valor"])
    por_lista: dict[str, int] = {}
    for c in resultado.coincidencias:
        por_lista[c.lista] = por_lista.get(c.lista, 0) + 1

    filas_resumen = [
        ("Clientes evaluados", resultado.clientes_evaluados),
        ("Designados en padron", resultado.designados_evaluados),
        ("Coincidencias totales", len(resultado.coincidencias)),
        ("Clientes con coincidencia", resultado.clientes_con_coincidencia),
        ("Casos escalados", sum(1 for c in casos if c.estado.value == "ESCALADO")),
        ("Umbral de revision", min((c.score for c in resultado.coincidencias), default=0)),
        ("Umbral de probable", umbral_probable),
    ]
    for lista, cantidad in sorted(por_lista.items()):
        filas_resumen.append((f"Coincidencias en {lista}", cantidad))

    for fila in filas_resumen:
        resumen.append(list(fila))
    for fila in resumen.iter_rows(min_row=2):
        for celda in fila:
            celda.font = _CUERPO
    _ajustar_anchos(resumen, [34, 18])

    # --- Hoja 2: coincidencias ---
    hoja = libro.create_sheet("Coincidencias")
    columnas = [
        "cliente_id", "lista", "id_origen", "designado", "matcheo_contra",
        "score", "criterio", "programas", "atenuantes", "prioridad",
    ]
    _escribir_encabezado(hoja, columnas)

    for c in resultado.coincidencias:
        prioridad = "PROBABLE" if c.score >= umbral_probable else "REVISAR"
        hoja.append([
            c.cliente_id, c.lista, c.id_origen, c.nombre_designado,
            c.nombre_matcheado, c.score, c.criterio,
            "; ".join(c.programas), "; ".join(c.atenuantes), prioridad,
        ])
        relleno = _PROBABLE if prioridad == "PROBABLE" else _REVISION
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
            celda.fill = relleno
    _ajustar_anchos(hoja, [14, 18, 12, 34, 30, 9, 12, 20, 40, 12])

    # --- Hoja 3: expediente ---
    exp = libro.create_sheet("Expediente")
    _escribir_encabezado(exp, ["caso_id", "cliente", "momento", "actor", "accion", "detalle"])
    for caso in casos:
        for ev in caso.evidencia:
            exp.append([
                caso.caso_id,
                caso.cliente.nombre,
                ev.momento.isoformat(timespec="seconds"),
                ev.actor,
                ev.accion,
                "; ".join(f"{k}={v}" for k, v in ev.detalle.items()),
            ])
    for fila in exp.iter_rows(min_row=2):
        for celda in fila:
            celda.font = _CUERPO
    _ajustar_anchos(exp, [12, 28, 22, 20, 22, 90])

    libro.save(ruta)
    return ruta


_NIVEL_RELLENO = {
    "ALTO": PatternFill("solid", start_color="F4B6B0"),
    "MEDIO": PatternFill("solid", start_color="FFE599"),
    "BAJO": PatternFill("solid", start_color="D9EAD3"),
}


def agregar_hojas_etapa2(
    ruta: str | Path,
    resoluciones: dict[str, "Resolucion"],
    evaluaciones: dict[str, "Evaluacion"],
    casos: list[Caso],
) -> Path:
    """Agrega las hojas de beneficiario final y riesgo al informe existente.

    Se escribe sobre el mismo libro que produjo la etapa 1 en lugar de emitir
    un segundo archivo: el analista recibe un informe, no una carpeta.
    """
    ruta = Path(ruta)
    libro = load_workbook(ruta)
    nombres = {c.cliente.cliente_id: c.cliente.nombre for c in casos}

    # --- Beneficiario final ---
    hoja = libro.create_sheet("Beneficiario final")
    _escribir_encabezado(hoja, [
        "cliente_id", "cliente", "beneficiario", "capital", "voto",
        "via", "detalle", "umbral", "niveles", "titularidad_opaca", "observaciones",
    ])

    for cliente_id, r in resoluciones.items():
        obs = "; ".join(r.observaciones)
        etiqueta = "EXCEPTUADA (oferta publica)" if r.exceptuada else "NO IDENTIFICADO"
        if not r.beneficiarios:
            hoja.append([
                cliente_id, nombres.get(cliente_id, ""), etiqueta,
                None, None, "", "", r.umbral_aplicado, r.profundidad_maxima,
                r.titularidad_opaca, obs,
            ])
        for b in r.beneficiarios:
            hoja.append([
                cliente_id, nombres.get(cliente_id, ""), b.nombre,
                b.capital, b.voto, b.via, b.detalle, r.umbral_aplicado,
                r.profundidad_maxima, r.titularidad_opaca, obs,
            ])

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.font = _CUERPO
        for idx in (4, 5, 8, 10):  # capital, voto, umbral, opaca
            fila[idx - 1].number_format = "0.00%"
    _ajustar_anchos(hoja, [12, 26, 28, 10, 10, 14, 46, 9, 9, 16, 50])

    # --- Riesgo ---
    hoja = libro.create_sheet("Riesgo")
    _escribir_encabezado(hoja, [
        "cliente_id", "cliente", "puntaje", "nivel", "regimen",
        "proxima_revision", "elevadores", "factores aplicados",
    ])

    orden = {"ALTO": 0, "MEDIO": 1, "BAJO": 2}
    for cliente_id, ev in sorted(
        evaluaciones.items(), key=lambda kv: (orden[kv[1].nivel], -kv[1].puntaje)
    ):
        hoja.append([
            cliente_id, nombres.get(cliente_id, ""), ev.puntaje, ev.nivel, ev.regimen,
            ev.proxima_revision().isoformat(), "; ".join(ev.elevadores),
            " | ".join(f"{f.codigo} (+{f.puntos:g}) {f.descripcion}" for f in ev.factores),
        ])
        relleno = _NIVEL_RELLENO.get(ev.nivel)
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
            if relleno:
                celda.fill = relleno
    _ajustar_anchos(hoja, [12, 26, 9, 9, 18, 16, 30, 110])

    libro.save(ruta)
    return ruta


COLUMNAS_OPERACIONES = [
    "cliente_id", "fecha", "monto", "sentido", "instrumento", "canal",
    "contraparte", "pais_contraparte", "referencia",
]

COLUMNAS_PERFILES = [
    "cliente_id", "monto_mensual", "operaciones_mensuales",
    "proporcion_efectivo", "paises", "origen_fondos", "proposito",
]


def _fecha_iso(crudo: str):
    from datetime import datetime
    texto = (crudo or "").strip()
    if not texto:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto.split(" ")[0], formato).date()
        except ValueError:
            continue
    return None


def _numero(crudo: str) -> float:
    """Acepta 1234.56, 1.234,56 y $ 1.234,56."""
    texto = (crudo or "").strip().replace("$", "").replace(" ", "")
    if not texto:
        return 0.0
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return 0.0


def leer_operaciones(ruta: str | Path) -> list[Operacion]:
    """Lee la operatoria. Una fila por operacion."""
    operaciones: list[Operacion] = []
    for fila in _leer_tabla(Path(ruta)):
        cliente_id = (fila.get("cliente_id") or "").strip()
        fecha = _fecha_iso(fila.get("fecha", ""))
        monto = abs(_numero(fila.get("monto", "")))
        if not cliente_id or fecha is None or monto <= 0:
            continue

        operaciones.append(Operacion(
            cliente_id=cliente_id,
            fecha=fecha,
            monto=monto,
            sentido=(fila.get("sentido") or "INGRESO").strip().upper() or "INGRESO",
            instrumento=(fila.get("instrumento") or "TRANSFERENCIA").strip().upper()
                        or "TRANSFERENCIA",
            canal=(fila.get("canal") or "ELECTRONICO").strip().upper() or "ELECTRONICO",
            contraparte=(fila.get("contraparte") or "").strip(),
            pais_contraparte=(fila.get("pais_contraparte") or "").strip(),
            referencia=(fila.get("referencia") or "").strip(),
        ))
    return operaciones


def leer_perfiles(ruta: str | Path) -> dict[str, Perfil]:
    """Lee los perfiles transaccionales declarados.

    Los paises esperados se normalizan a ISO al entrar, igual que en el resto
    del sistema. Si el perfil dice "Uruguay" y la operacion dice "UY", tienen
    que ser lo mismo.
    """
    perfiles: dict[str, Perfil] = {}
    for fila in _leer_tabla(Path(ruta)):
        cliente_id = (fila.get("cliente_id") or "").strip()
        if not cliente_id:
            continue

        crudos = [p.strip() for p in (fila.get("paises") or "").split(";") if p.strip()]
        paises = tuple(c for c in (iso(p) for p in crudos) if c)

        perfiles[cliente_id] = Perfil(
            cliente_id=cliente_id,
            monto_mensual=_numero(fila.get("monto_mensual", "")),
            operaciones_mensuales=int(_numero(fila.get("operaciones_mensuales", ""))),
            proporcion_efectivo=_fraccion(fila.get("proporcion_efectivo", "")),
            paises=paises,
            origen_fondos=(fila.get("origen_fondos") or "").strip(),
            proposito=(fila.get("proposito") or "").strip(),
        )
    return perfiles


_VENCIDO = PatternFill("solid", start_color="D96459")

_SEVERIDAD_RELLENO = {
    "ALTA": PatternFill("solid", start_color="F4B6B0"),
    "MEDIA": PatternFill("solid", start_color="FFE599"),
    "BAJA": PatternFill("solid", start_color="E8E8E8"),
}


def agregar_hoja_alertas(
    ruta: str | Path,
    alertas_por_cliente: dict[str, list["Alerta"]],
    casos: list[Caso],
    evaluaciones: dict[str, "Evaluacion"],
    perfiles: dict[str, Perfil],
) -> Path:
    """Escribe el registro de operaciones inusuales.

    Las columnas no son arbitrarias: replican el contenido minimo que los
    manuales del sector exigen para este registro. Las dos ultimas salen
    vacias porque las completa el analista al resolver la alerta, y
    prellenarlas seria fingir un analisis que no ocurrio.
    """
    ruta = Path(ruta)
    libro = load_workbook(ruta)
    nombres = {c.cliente.cliente_id: c.cliente.nombre for c in casos}

    hoja = libro.create_sheet("Alertas")
    columnas = [
        "alerta_id", "cliente_id", "cliente", "nivel_riesgo", "perfil_declarado",
        "tipo_inusualidad", "regimen", "severidad", "descripcion",
        "monto_involucrado", "operaciones", "metodologia", "generada",
        "vence", "estado_plazo", "resolucion", "medidas_adoptadas",
        "decision_final", "fecha_decision",
    ]
    _escribir_encabezado(hoja, columnas)

    orden = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    filas = [
        (cid, a)
        for cid, alertas in alertas_por_cliente.items()
        for a in alertas
    ]
    filas.sort(key=lambda par: (orden[par[1].severidad], -par[1].monto_involucrado))

    for cliente_id, a in filas:
        ev = evaluaciones.get(cliente_id)
        perfil = perfiles.get(cliente_id)
        resumen_perfil = (
            f"${perfil.monto_mensual:,.0f}/mes, {perfil.operaciones_mensuales} op."
            if perfil and perfil.declarado else "no declarado"
        )
        estado = ("PLAZO VENCIDO" if a.vencida
                  else f"{a.dias_restantes} dia(s)")
        hoja.append([
            a.identificador,
            cliente_id,
            nombres.get(cliente_id, ""),
            ev.nivel if ev else "",
            resumen_perfil,
            a.codigo,
            a.regimen.value,
            a.severidad,
            a.descripcion,
            a.monto_involucrado,
            a.detalle_operaciones,
            a.metodologia,
            a.generada.isoformat(),
            a.vence.isoformat(),
            estado,
            "",
            "",
            "",
            "",
        ])
        relleno = _VENCIDO if a.vencida else _SEVERIDAD_RELLENO.get(a.severidad)
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
            if relleno:
                celda.fill = relleno
        hoja.cell(row=hoja.max_row, column=10).number_format = "#,##0"

    if not filas:
        hoja.append(["sin alertas de monitoreo"] + [""] * (len(columnas) - 1))
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO

    _ajustar_anchos(hoja, [14, 12, 26, 9, 26, 26, 8, 10, 72, 16, 70, 52, 12, 12,
                           15, 14, 40, 60, 14])

    # Lista desplegable en resolucion. El ida y vuelta pasa por este campo, y
    # un "reportar" mal tipeado se lee como pendiente y la alerta desaparece
    # del circuito sin que nadie se entere.
    validacion = DataValidation(
        type="list", formula1='"REPORTAR,JUSTIFICADA"', allow_blank=True,
        showDropDown=False,
    )
    validacion.error = "Solo REPORTAR o JUSTIFICADA"
    validacion.errorTitle = "Resolucion invalida"
    hoja.add_data_validation(validacion)
    if hoja.max_row > 1:
        validacion.add(f"P2:P{hoja.max_row}")

    libro.save(ruta)
    return ruta


def agregar_hojas_etapa4(
    ruta: str | Path,
    congelamientos: list,
    exposicion,
) -> Path:
    """Agrega las hojas de congelamiento y exposicion sancionatoria.

    La hoja de congelamiento lleva una advertencia de reserva en la primera
    fila. No es un adorno: la norma obliga a abstenerse de informar al cliente
    los antecedentes de la resolucion, y un informe que circule sin esa marca
    es un aviso esperando ocurrir.
    """
    ruta = Path(ruta)
    libro = load_workbook(ruta)

    # --- Congelamiento ---
    hoja = libro.create_sheet("Congelamiento")
    aviso = ("CONFIDENCIAL. Art. 21 inc. c) y 22 Ley 25.246. Prohibido informar al "
             "cliente o a terceros los antecedentes de la medida.")
    hoja.append([aviso])
    hoja["A1"].font = Font(name=FUENTE, bold=True, color="9C0006", size=10)
    hoja["A1"].fill = PatternFill("solid", start_color="FFC7CE")

    columnas = ["cliente_id", "cliente", "regimen", "lista", "designado",
                "id_origen", "score", "criterio", "detectado", "norma", "paso",
                "obligacion", "cumplido"]
    hoja.append(columnas)
    for celda in hoja[2]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.freeze_panes = "A3"

    for c in congelamientos:
        for i, paso in enumerate(c.pasos, start=1):
            hoja.append([
                c.cliente_id, c.cliente, c.regimen.value, c.lista, c.designado,
                c.id_origen, c.score, c.criterio, c.detectado.isoformat(),
                c.norma, f"{i}. {paso.codigo}", paso.descripcion,
                "SI" if paso.cumplido else "",
            ])
            for celda in hoja[hoja.max_row]:
                celda.font = _CUERPO

    if not congelamientos:
        hoja.append(["sin obligaciones de congelamiento"] + [""] * (len(columnas) - 1))
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO

    _ajustar_anchos(hoja, [12, 26, 9, 16, 30, 12, 8, 12, 12, 44, 26, 90, 10])

    # --- Exposicion ---
    hoja = libro.create_sheet("Exposicion")
    _escribir_encabezado(hoja, ["concepto", "materia", "modulos", "monto", "fundamento"])

    for cargo in sorted(exposicion.cargos, key=lambda c: -c.monto):
        hoja.append([cargo.concepto, cargo.materia, cargo.modulos or "",
                     cargo.monto, cargo.fundamento])
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
        hoja.cell(row=hoja.max_row, column=4).number_format = "#,##0"

    hoja.append([])
    hoja.append(["TOTAL ESTIMADO", "", "", exposicion.total, ""])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(name=FUENTE, bold=True, size=10)
    hoja.cell(row=hoja.max_row, column=4).number_format = "#,##0"

    hoja.append([])
    hoja.append([
        f"Modulo ${exposicion.modulo_aplicado:,.0f} vigente desde "
        f"{exposicion.modulo_vigencia.isoformat()}. Estimacion de exposicion, no "
        f"calculo de multa: la sancion efectiva la determina la UIF en sumario."
    ])
    hoja[f"A{hoja.max_row}"].font = Font(name=FUENTE, italic=True, size=9)

    _ajustar_anchos(hoja, [46, 34, 10, 18, 56])

    libro.save(ruta)
    return ruta


COLUMNAS_DECISION = {
    "cliente_id", "tipo_inusualidad", "resolucion", "medidas_adoptadas",
    "decision_final", "fecha_decision",
}


def leer_decisiones(ruta: str | Path) -> dict:
    """Lee de vuelta lo que el analista resolvio en la hoja Alertas.

    El equipo de cumplimiento trabaja en Excel y ahi se queda. El sistema le
    entrega el informe con las columnas de decision vacias y lo lee cuando
    vuelve. Obligarlo a cargar las conclusiones en otra herramienta seria la
    solucion prolija que nadie usa.
    """
    from .ros import Decision, Resolucion

    libro = load_workbook(Path(ruta), read_only=True, data_only=True)
    if "Alertas" not in libro.sheetnames:
        libro.close()
        return {}

    hoja = libro["Alertas"]
    filas = hoja.iter_rows(values_only=True)
    encabezados = [str(c or "").strip() for c in next(filas)]
    indice = {nombre: i for i, nombre in enumerate(encabezados)}

    def valor(fila, nombre: str) -> str:
        i = indice.get(nombre)
        if i is None or i >= len(fila) or fila[i] is None:
            return ""
        return str(fila[i]).strip()

    decisiones: dict = {}
    for fila in filas:
        alerta_id = valor(fila, "alerta_id")
        cliente_id = valor(fila, "cliente_id")
        codigo = valor(fila, "tipo_inusualidad")
        if not alerta_id or not cliente_id:
            continue

        crudo = valor(fila, "resolucion").upper()
        try:
            resolucion = Resolucion(crudo) if crudo else Resolucion.PENDIENTE
        except ValueError:
            resolucion = Resolucion.PENDIENTE

        decisiones[alerta_id] = Decision(
            alerta_id=alerta_id,
            cliente_id=cliente_id,
            codigo_alerta=codigo,
            resolucion=resolucion,
            medidas=valor(fila, "medidas_adoptadas"),
            motivo=valor(fila, "decision_final"),
            fecha=_fecha_iso(valor(fila, "fecha_decision")),
        )

    libro.close()
    return decisiones


def agregar_hojas_etapa5(ruta: str | Path, resultado) -> Path:
    """Escribe los borradores de ROS y el registro de inusuales.

    Son dos salidas y las dos son obligatorias. La segunda es la que se
    olvida: las inusualidades resueltas sin reportar tambien tienen que
    quedar registradas con su analisis.
    """
    ruta = Path(ruta)
    libro = load_workbook(ruta)

    # --- Borradores de ROS ---
    hoja = libro.create_sheet("Borradores ROS")
    aviso = ("CONFIDENCIAL. Art. 21 inc. c) y 22 Ley 25.246. Insumo para cargar en el "
             "SRO+ de la UIF; no constituye un reporte presentado.")
    hoja.append([aviso])
    hoja["A1"].font = Font(name=FUENTE, bold=True, color="9C0006", size=10)
    hoja["A1"].fill = PatternFill("solid", start_color="FFC7CE")

    columnas = ["cliente_id", "cliente", "tipo", "documento", "regimen",
                "nivel_riesgo", "monto", "vence", "estado", "faltantes",
                "observaciones", "fundamento"]
    hoja.append(columnas)
    for celda in hoja[2]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.freeze_panes = "A3"

    for b in resultado.borradores:
        documento = b.cliente.documentos[0].clave() if b.cliente.documentos else ""
        estado = ("FUERA DE PLAZO" if b.fuera_de_plazo
                  else "PRESENTABLE" if b.completo else "INCOMPLETO")
        hoja.append([
            b.cliente.cliente_id, b.cliente.nombre, b.cliente.tipo, documento,
            b.regimen.value, b.nivel_riesgo, b.monto, b.vence.isoformat(),
            estado, "; ".join(b.faltantes()), "; ".join(b.observaciones()),
            b.fundamento(),
        ])
        relleno = (_VENCIDO if b.fuera_de_plazo
                   else None if b.completo else _SEVERIDAD_RELLENO["MEDIA"])
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if relleno:
                celda.fill = relleno
        hoja.cell(row=hoja.max_row, column=7).number_format = "#,##0"

    if not resultado.borradores:
        hoja.append(["sin operaciones resueltas para reportar"] + [""] * (len(columnas) - 1))
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO

    _ajustar_anchos(hoja, [12, 26, 10, 18, 9, 12, 16, 12, 16, 40, 56, 110])

    # --- Registro de inusuales no reportadas ---
    hoja = libro.create_sheet("Inusuales justificadas")
    _escribir_encabezado(hoja, [
        "cliente_id", "cliente", "nivel_riesgo", "tipo_inusualidad",
        "descripcion", "monto", "generada", "medidas_adoptadas",
        "decision_final", "fecha_decision", "documentada",
    ])

    for j in resultado.justificadas:
        hoja.append([
            j.cliente_id, j.cliente, j.nivel_riesgo, j.alerta.codigo,
            j.alerta.descripcion, j.alerta.monto_involucrado,
            j.alerta.generada.isoformat(), j.decision.medidas, j.decision.motivo,
            j.decision.fecha.isoformat() if j.decision.fecha else "",
            "SI" if j.documentada else "NO",
        ])
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO
            if not j.documentada:
                celda.fill = _SEVERIDAD_RELLENO["MEDIA"]
        hoja.cell(row=hoja.max_row, column=6).number_format = "#,##0"

    if not resultado.justificadas:
        hoja.append(["sin inusualidades justificadas"] + [""] * 10)
        for celda in hoja[hoja.max_row]:
            celda.font = _CUERPO

    _ajustar_anchos(hoja, [12, 26, 10, 26, 72, 16, 12, 44, 60, 14, 12])

    libro.save(ruta)
    return ruta
